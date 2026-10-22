"""
export_excel.py - writes the database out as an Excel workbook with charts.

    python export_excel.py

Creates exports/Montana_Deer_Harvest.xlsx with four sheets:
  Summary         statewide totals per year (formulas) + line chart
  Top Districts   best-odds districts for the latest surveyed year + bar charts
  Trends          every district's harvest by year, colour-shaded, with % change
  Data            every row from the database, as a filterable Excel table

Uses openpyxl only (pure Python), so it runs even with Smart App Control on.
"""

from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from database import connect

ROOT = Path(__file__).parent
OUT = ROOT / "exports" / "Montana_Deer_Harvest.xlsx"
SPECIES = ["mule deer", "whitetail"]
TOP_N = 25

FONT = "Arial"
H1 = Font(name=FONT, size=16, bold=True, color="1F2A22")
H2 = Font(name=FONT, size=12, bold=True, color="2F5D3A")
BOLD = Font(name=FONT, bold=True)
BODY = Font(name=FONT)
NOTE = Font(name=FONT, italic=True, size=9, color="5D6A60")
HEAD_FILL = PatternFill("solid", fgColor="2F5D3A")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF")


def load():
    conn = connect()
    rows = [dict(r) for r in conn.execute("""
        SELECT district, species, year, hunters, hunter_days, harvest
        FROM harvest_stats WHERE harvest IS NOT NULL ORDER BY species, district, year""")]
    conn.close()
    return rows


def header(ws, row, labels, start_col=1):
    for i, text in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=text)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(horizontal="center", wrap_text=True)


def widths(ws, sizes):
    for col, w in sizes.items():
        ws.column_dimensions[col].width = w


def data_sheet(wb, rows):
    ws = wb.create_sheet("Data")
    labels = ["District", "Region", "Species", "Year", "Hunters", "Hunter days", "Harvest",
              "Per hunter", "Days per animal"]
    header(ws, 1, labels)
    for i, r in enumerate(rows, start=2):
        ws.append([r["district"], int(r["district"][0]), r["species"], r["year"],
                   r["hunters"], r["hunter_days"], r["harvest"],
                   f'=IF(E{i}>0,G{i}/E{i},"")', f'=IF(AND(F{i}>0,G{i}>0),F{i}/G{i},"")'])
        for c in ws[i]:
            c.font = BODY
        ws[f"E{i}"].number_format = ws[f"F{i}"].number_format = ws[f"G{i}"].number_format = "#,##0"
        ws[f"H{i}"].number_format = "0%"
        ws[f"I{i}"].number_format = "0.0"
    last = len(rows) + 1
    tab = Table(displayName="Harvest", ref=f"A1:I{last}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium7", showRowStripes=True)
    ws.add_table(tab)
    ws.freeze_panes = "A2"
    widths(ws, {"A": 10, "B": 8, "C": 12, "D": 8, "E": 10, "F": 12, "G": 10, "H": 11, "I": 10})
    return last


def summary_sheet(wb, years, last):
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Montana deer harvest, statewide"
    ws["A1"].font = H1
    ws["A2"] = "Totals are formulas over the Data sheet, so they update if you edit it."
    ws["A2"].font = NOTE
    header(ws, 4, ["Year", "Mule deer", "Whitetail", "All deer", "Change vs prior year"])
    for i, y in enumerate(years, start=5):
        ws[f"A{i}"] = y
        ws[f"B{i}"] = f'=SUMIFS(Data!$G$2:$G${last},Data!$C$2:$C${last},"mule deer",Data!$D$2:$D${last},A{i})'
        ws[f"C{i}"] = f'=SUMIFS(Data!$G$2:$G${last},Data!$C$2:$C${last},"whitetail",Data!$D$2:$D${last},A{i})'
        ws[f"D{i}"] = f"=B{i}+C{i}"
        ws[f"E{i}"] = f'=IF(D{i - 1}>0,D{i}/D{i - 1}-1,"")' if i > 5 else "–"
        for col in "ABCDE":
            ws[f"{col}{i}"].font = BODY
        for col in "BCD":
            ws[f"{col}{i}"].number_format = "#,##0"
        ws[f"E{i}"].number_format = "+0%;-0%;0%"
        ws[f"E{i}"].alignment = Alignment(horizontal="right")
    end = 4 + len(years)

    chart = LineChart()
    chart.title = "Deer harvested statewide"
    chart.y_axis.title = "Animals"
    chart.y_axis.numFmt = "#,##0"
    chart.height, chart.width = 8, 16
    chart.add_data(Reference(ws, min_col=2, max_col=3, min_row=4, max_row=end), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=5, max_row=end))
    for s, color in zip(chart.series, ["A0522D", "2F5D3A"]):
        s.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.width = 28000
        s.smooth = False
    ws.add_chart(chart, "G4")

    ws[f"A{end + 2}"] = ("Source: Montana FWP harvest estimates (myfwp.mt.gov). FWP counts deer hunters per "
                         "district, not per species. Some years (e.g. 2022, 2024) have no hunter survey.")
    ws[f"A{end + 2}"].font = NOTE
    widths(ws, {"A": 10, "B": 13, "C": 13, "D": 13, "E": 13})


def top_sheet(wb, rows):
    ws = wb.create_sheet("Top Districts")
    surveyed = sorted({r["year"] for r in rows if r["hunters"]})
    year = surveyed[-1]
    ws["A1"] = f"Best-odds districts, {year}"
    ws["A1"].font = H1
    ws["A2"] = (f"Ranked by harvest per deer hunter, districts with 100+ hunters. {year} is the latest year "
                "FWP surveyed hunter numbers.")
    ws["A2"].font = NOTE

    col = 1
    for sp in SPECIES:
        pick = [r for r in rows if r["species"] == sp and r["year"] == year and (r["hunters"] or 0) >= 100]
        pick.sort(key=lambda r: r["harvest"] / r["hunters"], reverse=True)
        pick = pick[:TOP_N]
        c0 = get_column_letter(col)
        ws[f"{c0}4"] = sp.title()
        ws[f"{c0}4"].font = H2
        header(ws, 5, ["#", "District", "Hunters", "Harvest", "Per hunter", "Days per animal"], start_col=col)
        for i, r in enumerate(pick, start=6):
            L = [get_column_letter(col + k) for k in range(6)]
            ws[f"{L[0]}{i}"] = i - 5
            ws[f"{L[1]}{i}"] = r["district"]
            ws[f"{L[2]}{i}"] = r["hunters"]
            ws[f"{L[3]}{i}"] = r["harvest"]
            ws[f"{L[4]}{i}"] = f"={L[3]}{i}/{L[2]}{i}"
            ws[f"{L[5]}{i}"] = r["hunter_days"] / r["harvest"] if r["hunter_days"] and r["harvest"] else None
            for k, fmt in enumerate(["0", "@", "#,##0", "#,##0", "0%", "0.0"]):
                ws[f"{L[k]}{i}"].number_format = fmt
                ws[f"{L[k]}{i}"].font = BODY
        for k, w in enumerate([5, 10, 10, 10, 11, 10]):
            ws.column_dimensions[get_column_letter(col + k)].width = w

        bar = BarChart()
        bar.type = "bar"
        bar.title = f"{sp.title()}: top 10 by harvest per hunter"
        bar.y_axis.numFmt = "0%"
        bar.y_axis.majorGridlines = None
        bar.legend = None
        bar.height, bar.width = 8, 12
        n = min(10, len(pick))
        bar.add_data(Reference(ws, min_col=col + 4, min_row=5, max_row=5 + n), titles_from_data=True)
        bar.set_categories(Reference(ws, min_col=col + 1, min_row=6, max_row=5 + n))
        bar.x_axis.scaling.orientation = "maxMin"      # #1 at the top
        bar.series[0].graphicalProperties.solidFill = "A0522D" if sp == "mule deer" else "2F5D3A"
        ws.add_chart(bar, f"{c0}{8 + TOP_N}")
        col += 7
    ws.freeze_panes = "A6"


def trends_sheet(wb, rows, years):
    ws = wb.create_sheet("Trends")
    ws["A1"] = "Harvest by district and year"
    ws["A1"].font = H1
    ws["A2"] = "Darker green = more animals taken (shaded per species). Change compares the first and latest year."
    ws["A2"].font = NOTE

    table = defaultdict(dict)
    for r in rows:
        table[(r["species"], r["district"])][r["year"]] = r["harvest"]

    row = 4
    for sp in SPECIES:
        ws.cell(row=row, column=1, value=sp.title()).font = H2
        row += 1
        header(ws, row, ["District", "Region"] + [str(y) for y in years] + ["Change"])
        first = row + 1
        for district in sorted(d for s, d in table if s == sp):
            row += 1
            vals = table[(sp, district)]
            ws.cell(row=row, column=1, value=district).font = BOLD
            ws.cell(row=row, column=2, value=int(district[0])).font = BODY
            for k, y in enumerate(years):
                c = ws.cell(row=row, column=3 + k, value=vals.get(y))
                c.number_format, c.font = "#,##0", BODY
            a, b = get_column_letter(3), get_column_letter(2 + len(years))
            c = ws.cell(row=row, column=3 + len(years),
                        value=f'=IF(AND(ISNUMBER({a}{row}),{a}{row}>0,ISNUMBER({b}{row})),{b}{row}/{a}{row}-1,"")')
            c.number_format, c.font = "+0%;-0%;0%", BODY
        rng = f"C{first}:{get_column_letter(2 + len(years))}{row}"
        ws.conditional_formatting.add(rng, ColorScaleRule(start_type="min", start_color="FFFDF8",
                                                          end_type="percentile", end_value=95, end_color="5E9C6A"))
        chg = f"{get_column_letter(3 + len(years))}{first}:{get_column_letter(3 + len(years))}{row}"
        ws.conditional_formatting.add(chg, ColorScaleRule(start_type="num", start_value=-0.6, start_color="E8A397",
                                                          mid_type="num", mid_value=0, mid_color="FFFFFF",
                                                          end_type="num", end_value=0.6, end_color="9CCB9F"))
        row += 3
    widths(ws, {"A": 10, "B": 8})
    for k in range(len(years) + 1):
        ws.column_dimensions[get_column_letter(3 + k)].width = 10
    ws.freeze_panes = "C4"


def build():
    rows = load()
    if not rows:
        raise SystemExit("No harvest data in hunting.db yet. Run run_all.py first.")
    years = sorted({r["year"] for r in rows})
    wb = Workbook()
    last = data_sheet(wb, rows)
    summary_sheet(wb, years, last)
    top_sheet(wb, rows)
    trends_sheet(wb, rows, years)
    wb.move_sheet("Data", offset=3)
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(f"Saved {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    build()
